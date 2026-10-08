import os
import re
import io
import streamlit as st
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# --- BUSINESS LOGIC CORE (FROM YOUR CODE) ---
FIELDS = [
    "No", "ID", "Name", "Alias", "Place/Date of Birth",
    "Nationality / Country of Origin", "Address", "Remarks", "Source File"
]

def clean_value(value):
    if value is None:
        return ""
    value = re.sub(r"\s*\n\s*", " ", value)
    value = re.sub(r"[ \t]{2,}", " ", value)
    return value.strip(" ;\n")

def get_section(text, start_marker, end_markers):
    start = text.find(start_marker)
    if start < 0:
        return ""
    end = len(text)
    for marker in end_markers:
        pos = text.find(marker, start + len(start_marker))
        if pos >= 0:
            end = min(end, pos)
    return text[start:end]

def parse_section(section_text, source_file):
    if not section_text:
        return []

    record_re = re.compile(r"(?m)^(?:(\d+)\s+)?(\d+)\.\s+Nama\s*(?::\s*|(?=[A-ZÀ-ÖØ-Ý]))")
    starts = list(record_re.finditer(section_text))
    if not starts:
        return []

    label_re = re.compile(
        r"(?mi)^\s*(Nama\s+alias|Tempat\s+tanggal\s+lahir|"
        r"Kewarganegaraan|Asal\s+negara|Alamat|Keterangan)\s*:?\s*"
    )

    rows = []
    # Clean programmatic counter that acts as our ultimate source of truth
    expected_counter = 1 

    for idx, match in enumerate(starts):
        block_end = starts[idx + 1].start() if idx + 1 < len(starts) else len(section_text)
        block = section_text[match.end():block_end]

        fields = {key: "" for key in FIELDS}
        
        # Extract the raw string digits found by regex
        raw_no = (match.group(1) + match.group(2)) if match.group(1) else match.group(2)
        
        # Fix conjoined page-number artifacts (e.g., '1426' when expected counter is around 26)
        if len(raw_no) >= 3:
            # If the number ends with our expected sequence value, isolate it
            expected_str = str(expected_counter)
            if raw_no.endswith(expected_str):
                fields["No"] = expected_str
            else:
                # If it's a completely scrambled artifact, trust our sequential counter
                fields["No"] = expected_str
        else:
            fields["No"] = raw_no

        fields["Source File"] = source_file

        labels = list(label_re.finditer(block))
        if labels:
            fields["Name"] = clean_value(block[:labels[0].start()]) # Fixed the bug here!
        else:
            fields["Name"] = clean_value(block)

        for j, label_match in enumerate(labels):
            raw_label = re.sub(r"\s+", " ", label_match.group(1)).lower()
            label_map = {
                "nama alias": "Alias",
                "tempat tanggal lahir": "Place/Date of Birth",
                "kewarganegaraan": "Nationality / Country of Origin",
                "asal negara": "Nationality / Country of Origin",
                "alamat": "Address",
                "keterangan": "Remarks",
            }
            key = label_map.get(raw_label)
            if not key:
                continue
            value_end = labels[j + 1].start() if j + 1 < len(labels) else len(block)
            fields[key] = clean_value(block[label_match.end():value_end])

        id_match = re.search(r"\(([^()]+)\)\s*$", fields["Name"])
        if id_match:
            fields["ID"] = id_match.group(1).strip()
            fields["Name"] = fields["Name"][:id_match.start()].strip(" ;")

        for key in FIELDS:
            fields[key] = clean_value(fields[key])

        rows.append(fields)
        expected_counter += 1  # Increment to track what the next row number should mathematically look like

    return rows


def parse_streamlit_pdf(uploaded_file):
    parts = []
    # Open from Streamlit memory stream directly
    with pdfplumber.open(uploaded_file) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    text = "\n".join(parts)

    individuals_section = get_section(text, "II. INDIVIDU:", ["III. ENTITAS"])
    entities_section = get_section(
        text, "III. ENTITAS:", 
        ["IV KETERANGAN", "IV. KETERANGAN", "IV KETERANGAN:", "IV. KETERANGAN:"]
    )

    individuals = parse_section(individuals_section, uploaded_file.name)
    entities = parse_section(entities_section, uploaded_file.name)

    for row in individuals:
        row["Section"] = "INDIVIDU"
    for row in entities:
        row["Section"] = "ENTITAS"

    return individuals, entities

def create_excel_buffer(all_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "DTTOT_ALL"

    headers = ["Source File", "Section", "No", "ID", "Name", "Alias",
               "Place/Date of Birth", "Nationality / Country of Origin",
               "Address", "Remarks"]
    ws.append(headers)

    for row in all_rows:
        ws.append([
            row.get("Source File", ""), row.get("Section", ""), row.get("No", ""),
            row.get("ID", ""), row.get("Name", ""), row.get("Alias", ""),
            row.get("Place/Date of Birth", ""), row.get("Nationality / Country of Origin", ""),
            row.get("Address", ""), row.get("Remarks", "")
        ])

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 24

    widths = {
        "A": 32, "B": 14, "C": 8, "D": 16, "E": 38,
        "F": 55, "G": 35, "H": 32, "I": 55, "J": 90
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    if ws.max_row >= 2:
        ref = f"A1:J{ws.max_row}"
        table = Table(displayName="DTTOTData", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True,
                               showColumnStripes=False)
        table.tableStyleInfo = style
        ws.add_table(table)

    summary = wb.create_sheet("Summary")
    summary.append(["Item", "Count"])
    individual_count = sum(1 for r in all_rows if r.get("Section") == "INDIVIDU")
    entity_count = sum(1 for r in all_rows if r.get("Section") == "ENTITAS")
    source_count = len({r.get("Source File") for r in all_rows})
    
    summary.append(["Source PDF files", source_count])
    summary.append(["Individual records", individual_count])
    summary.append(["Entity records", entity_count])
    summary.append(["Total records", len(all_rows)])
    
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 16

    # Save to binary memory stream instead of physical hard drive
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer


# --- STREAMLIT USER INTERFACE ---
st.set_page_config(page_title="DTTOT PDF Converter", page_icon="📄", layout="wide")

st.title("📄 DTTOT Batch PDF → Excel Converter")
st.markdown("Select multiple DTTOT PDF files. All Section II and III records will be merged into one download-ready spreadsheet.")

uploaded_files = st.file_uploader(
    "Drop your DTTOT PDF files here:", 
    type=["pdf"], 
    accept_multiple_files=True
)

if uploaded_files:
    st.info(f"📁 {len(uploaded_files)} file(s) selected. Ready to convert.")
    
    if st.button("🚀 Run Conversion Process", type="primary"):
        all_rows = []
        errors = []
        
        # UI Log Window Container
        log_container = st.expander("📝 Live Conversion Log", expanded=True)
        
        with st.spinner("Parsing data from structural layers..."):
            for idx, uploaded_file in enumerate(uploaded_files, 1):
                log_container.write(f"Reading {idx}/{len(uploaded_files)}: **{uploaded_file.name}**")
                try:
                    individuals, entities = parse_streamlit_pdf(uploaded_file)
                    all_rows.extend(individuals)
                    all_rows.extend(entities)
                    log_container.write(f"  → Found {len(individuals)} INDIVIDU + {len(entities)} ENTITAS")
                except Exception as exc:
                    errors.append((uploaded_file.name, str(exc)))
                    log_container.error(f"  ⚠️ Error handling file: {exc}")

        if not all_rows:
            st.error("No valid DTTOT records could be matched or extracted from the files.")
        else:
            # Metrics Visual Representation
            st.success("🎉 Structural parsing successful!")
            
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Processed Files", len(uploaded_files))
            m2.metric("Individual Entries", sum(1 for r in all_rows if r.get("Section") == "INDIVIDU"))
            m3.metric("Entity Entries", sum(1 for r in all_rows if r.get("Section") == "ENTITAS"))
            m4.metric("Total Records", len(all_rows))
            
            if errors:
                st.warning(f"Note: {len(errors)} file(s) encountered compilation snags. Review the live log layout above.")

            # Generate workbook in memory
            excel_data = create_excel_buffer(all_rows)
            
            # Big Interactive Download trigger
            st.download_button(
                label="📥 Click Here to Download Combined Excel Workbook",
                data=excel_data,
                file_name="DTTOT_Combined_Output.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
